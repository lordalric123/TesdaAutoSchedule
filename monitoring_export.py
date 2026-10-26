"""Simplified Assessment Monitoring Excel export."""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

from date_rules import parse_iso_date


def compact_day_span(start: date, end: date) -> str:
    if start == end:
        return str(start.day)
    if start.month == end.month and start.year == end.year:
        return f"{start.day}-{end.day}"
    if start.year == end.year:
        return f"{start.month}/{start.day}-{end.month}/{end.day}"
    return f"{start.month}/{start.day}/{start.year}-{end.month}/{end.day}/{end.year}"


def compact_assessors(assessors: list[dict]) -> str:
    if not assessors:
        return ""
    labels = [
        f"{a.get('name', '')} ({'Region' if a.get('assessor_type') == 'region' else 'Province'})"
        for a in assessors
        if a.get("name")
    ]
    return "\n".join(labels)


def compact_approved_dates(iso_dates: list[str]) -> str:
    if not iso_dates:
        return ""
    parsed = sorted({parse_iso_date(value) for value in iso_dates})
    first, last = parsed[0], parsed[-1]
    if first == last:
        return f"{first.month}/{first.day}/{first.year}"
    if first.month == last.month and first.year == last.year:
        return f"{first.month}/{first.day}-{last.day}/{last.year}"
    if first.year == last.year:
        return f"{first.month}/{first.day}-{last.month}/{last.day}, {last.year}"
    return f"{first.month}/{first.day}/{first.year}-{last.month}/{last.day}/{last.year}"


def _write_monitoring_sheet(sheet, month_label: str, items: list[dict]) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LEGAL
    sheet.page_margins.left = 0.4
    sheet.page_margins.right = 0.4

    title_fill = PatternFill("solid", fgColor="1F4E79")
    header_fill = PatternFill("solid", fgColor="FFC000")
    zebra = PatternFill("solid", fgColor="FFF2CC")
    thin = Border(
        left=Side(style="thin", color="7F7F7F"),
        right=Side(style="thin", color="7F7F7F"),
        top=Side(style="thin", color="7F7F7F"),
        bottom=Side(style="thin", color="7F7F7F"),
    )
    title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    header_font = Font(name="Calibri", size=11, bold=True)
    body_font = Font(name="Calibri", size=11)
    wrap = Alignment(wrap_text=True, vertical="center")
    center = Alignment(wrap_text=True, vertical="center", horizontal="center")

    sheet.merge_cells("A1:H1")
    title = sheet["A1"]
    title.value = f"ASSESSMENT SCHEDULES FOR {month_label.upper()}"
    title.fill = title_fill
    title.font = title_font
    title.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 28

    headers = [
        "APPROVED DATES",
        "ASSESSMENT CENTERS",
        "QUALIFICATIONS",
        "ASSESSMENT DATE",
        "NUMBER OF PAX",
        "ASSESSOR",
        "TESDA REPRESENTATIVE",
        "REMARKS",
    ]
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(2, col, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin
    sheet.row_dimensions[2].height = 22
    sheet.auto_filter.ref = f"A2:H{max(3, 2 + len(items))}"
    sheet.freeze_panes = "A3"

    for index, item in enumerate(items):
        row = 3 + index
        start = parse_iso_date(item["start_date"])
        end = parse_iso_date(item["end_date"])
        assessor_text = compact_assessors(item.get("assessors") or [])
        values = [
            compact_approved_dates(item.get("approved_date_list") or []),
            item.get("assessment_center") or "",
            item.get("qualification") or "",
            compact_day_span(start, end),
            item.get("pax") or "",
            assessor_text,
            item.get("tesda_representative") or "",
            item.get("remarks") or "",
        ]
        fill = zebra if index % 2 else None
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row, col, value)
            cell.font = body_font
            cell.alignment = center if col in {1, 4, 5} else wrap
            cell.border = thin
            if fill:
                cell.fill = fill
        line_count = max(1, assessor_text.count("\n") + 1)
        sheet.row_dimensions[row].height = max(32, 16 * line_count + 16)

    widths = [18, 42, 36, 16, 14, 28, 26, 30]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def build_monitoring_workbook(year: int, month: int, items: list[dict]) -> BytesIO:
    month_label = date(year, month, 1).strftime("%B %Y")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = month_label[:31]
    _write_monitoring_sheet(sheet, month_label, items)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def iter_months(start_year: int, start_month: int, end_year: int, end_month: int):
    """Yield (year, month) tuples from the start month through the end month, inclusive."""
    year, month = start_year, start_month
    end_index = end_year * 12 + (end_month - 1)
    while year * 12 + (month - 1) <= end_index:
        yield year, month
        month += 1
        if month > 12:
            month = 1
            year += 1


def _month_bounds(year: int, month: int) -> tuple[date, date]:
    month_start = date(year, month, 1)
    month_end = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)
    return month_start, month_end


def bucket_assessments_by_month(
    items: list[dict], months: list[tuple[int, int]]
) -> dict[tuple[int, int], list[dict]]:
    """Place each assessment into exactly one month in `months`: its own start month if that
    month is in range, otherwise the earliest month in range it overlaps. This keeps assessments
    that span a month boundary (e.g. Sept 29 - Oct 2) from being duplicated across both sheets.
    """
    buckets: dict[tuple[int, int], list[dict]] = {key: [] for key in months}
    month_set = set(months)
    for item in items:
        start = parse_iso_date(item["start_date"])
        end = parse_iso_date(item["end_date"])
        start_key = (start.year, start.month)
        if start_key in month_set:
            buckets[start_key].append(item)
            continue
        placed = False
        for key in months:
            month_start, month_end = _month_bounds(*key)
            if start < month_end and end >= month_start:
                buckets[key].append(item)
                placed = True
                break
        if not placed:
            # Assessment doesn't actually overlap the exported range; skip it.
            continue
    return buckets


def build_monitoring_workbook_range(
    start_year: int, start_month: int, end_year: int, end_month: int, items: list[dict]
) -> BytesIO:
    months = list(iter_months(start_year, start_month, end_year, end_month))
    buckets = bucket_assessments_by_month(items, months)

    workbook = Workbook()
    workbook.remove(workbook.active)
    used_titles: set[str] = set()
    for key in months:
        year, month = key
        month_label = date(year, month, 1).strftime("%B %Y")
        title = month_label[:31]
        base_title, suffix = title, 2
        while title in used_titles:
            title = f"{base_title[:28]} ({suffix})"
            suffix += 1
        used_titles.add(title)
        sheet = workbook.create_sheet(title=title)
        month_items = sorted(buckets[key], key=lambda x: x["start_date"])
        _write_monitoring_sheet(sheet, month_label, month_items)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def build_assessor_history_workbook(
    assessment_log: list[dict], assessor_type: str = "", group_by: str = "month"
) -> BytesIO:
    """One worksheet per month (or, with group_by="year", one worksheet per
    calendar year), laid out like the official TESDA-OP-CO-05-F35 "Assignment
    of Assessors" form: a title block, then one QUALIFICATION TITLE /
    PROVINCE (or REGION) table per qualification active that period, listing
    each assessor assignment underneath. Every qualification is always
    included — callers should not pre-filter by qualification, only by
    assessor type and (optionally) a date range.
    """
    group_by = "year" if group_by == "year" else "month"
    workbook = Workbook()
    workbook.remove(workbook.active)

    province_label = "REGION" if assessor_type == "region" else "PROVINCE"
    # Pre-filled for this office's own exports; only applies to the PROVINCE box —
    # a region-scoped export leaves the REGION box blank since this isn't a region name.
    province_value = "" if assessor_type == "region" else "Southern Leyte"

    # Font family/sizes matched to the official TESDA-OP-CO-05-F35 "Assignment of
    # Assessors" form: Arial throughout, bold caps for the title and field labels,
    # plain body text for the filled-in values.
    thin = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    form_code_font = Font(name="Arial", size=9)
    title_font = Font(name="Arial", size=16, bold=True)
    month_font = Font(name="Arial", size=11)
    label_font = Font(name="Arial", size=11, bold=True)
    header_font = Font(name="Arial", size=11, bold=True)
    body_font = Font(name="Arial", size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")

    # Each of the 3 visible columns (NAME OF ASSESSOR, ASSESSMENT CENTER, DATE OF
    # ASSESSMENT) is actually 4 equal-width sub-columns under the hood, so the
    # QUALIFICATION TITLE / PROVINCE row above the table can split at the exact
    # fractions of the official form: QUALIFICATION TITLE spans the first 3/4 of
    # NAME OF ASSESSOR; its value spans the remaining 1/4 of NAME OF ASSESSOR plus
    # the first 3/4 of ASSESSMENT CENTER; PROVINCE/REGION (label, then value) takes
    # what's left — the last 1/4 of ASSESSMENT CENTER plus all of DATE OF ASSESSMENT.
    headers = ["NAME OF ASSESSOR", "ASSESSMENT CENTER", "DATE OF ASSESSMENT"]
    # Generous enough that typical names/center names fit on one line; any longer
    # value still wraps (never gets cut off) and the row height is grown to match.
    base_widths = [28, 50, 24]
    SUB = 4  # sub-columns per visible column, so quarters land exactly on a column edge
    widths = [w / SUB for w in base_widths for _ in range(SUB)]
    ncols = len(widths)
    name_cols = (1, SUB)  # columns 1-4
    center_cols = (SUB + 1, SUB * 2)  # columns 5-8
    date_cols = (SUB * 2 + 1, SUB * 3)  # columns 9-12
    date_col_width = base_widths[2]
    QUAL_SPACER_ROWS = 3  # blank rows of breathing room between qualification tables

    def lines_needed(text, col_width: float, bold: bool = False) -> int:
        """Approximate how many wrapped lines `text` will take at `col_width`,
        so the row can be made tall enough that nothing gets clipped. Calibrated
        against actual Excel/LibreOffice wrapping at 11pt Arial: roughly 0.8
        characters fit per column-width unit (bold text runs a bit wider)."""
        if not text:
            return 1
        chars_per_line = max(6, int(col_width * (0.72 if bold else 0.8)))
        total = 0
        for segment in str(text).split("\n"):
            seg_len = len(segment) or 1
            total += max(1, -(-seg_len // chars_per_line))  # ceil division
        return max(1, total)

    def autofit_row_height(
        values: list, col_widths: list[float], per_line: float = 16, pad: float = 8, bold: bool = False
    ) -> float:
        lines = max(lines_needed(v, w, bold=bold) for v, w in zip(values, col_widths))
        return max(20.0, min(lines * per_line + pad, 160.0))

    groups: dict[tuple[int, ...], list[dict]] = {}
    for item in assessment_log:
        d = parse_iso_date(item["start_date"])
        key = (d.year,) if group_by == "year" else (d.year, d.month)
        groups.setdefault(key, []).append(item)

    def write_sheet_head(sheet, period_text: str) -> int:
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "portrait"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        sheet.page_margins.left = 0.5
        sheet.page_margins.right = 0.5
        sheet.page_margins.top = 0.6
        sheet.page_margins.bottom = 0.6
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width

        row = 1
        code_cell = sheet.cell(row, ncols, "TESDA-OP-CO-05-F35")
        code_cell.font = form_code_font
        code_cell.alignment = right
        row += 1
        rev_cell = sheet.cell(row, ncols, "Rev.No.00-03/08/17")
        rev_cell.font = form_code_font
        rev_cell.alignment = right
        row += 1

        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        title_cell = sheet.cell(row, 1, "ASSIGNMENT OF ASSESSORS")
        title_cell.font = title_font
        title_cell.alignment = center
        sheet.row_dimensions[row].height = 24
        row += 1

        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        month_cell = sheet.cell(row, 1, period_text)
        month_cell.font = month_font
        month_cell.alignment = center
        row += 2
        return row

    if not groups:
        sheet = workbook.create_sheet("Assessor History"[:31])
        row = write_sheet_head(sheet, "For the year —" if group_by == "year" else "For the month of —")
        sheet.cell(row, 1, "No recorded assessments.").font = body_font
    else:
        used_names: set[str] = set()
        for key in sorted(groups.keys()):
            if group_by == "year":
                (year,) = key
                period_text = f"For the year {year}"
                base_name = str(year)
            else:
                year, month = key
                month_date = date(year, month, 1)
                period_text = f"For the month of {month_date.strftime('%B %Y')}"
                base_name = month_date.strftime("%b %Y")
            sheet_name = base_name
            suffix = 1
            while sheet_name in used_names:
                suffix += 1
                sheet_name = f"{base_name} ({suffix})"
            used_names.add(sheet_name)
            sheet = workbook.create_sheet(sheet_name)
            row = write_sheet_head(sheet, period_text)

            # Printable area per page, in points, so a qualification table that
            # won't fully fit on the current page is pushed to a fresh one
            # instead of being cut mid-table. A4 portrait is 11.69in tall;
            # subtract the sheet's own top/bottom margins, then shave a
            # safety margin for the approximate row-height math above and for
            # renderer differences between Excel and LibreOffice.
            usable_height_pt = (
                11.69 - sheet.page_margins.top - sheet.page_margins.bottom
            ) * 72 * 0.9
            # The header block (form code, title, "For the month/year of…")
            # already occupies the top of the first page.
            page_used_pt = sum((sheet.row_dimensions[r].height or 15) for r in range(1, row))

            grouped: dict[str, list[dict]] = {}
            for item in groups[key]:
                grouped.setdefault(item.get("qualification") or "Unspecified Qualification", []).append(item)
            for entries in grouped.values():
                entries.sort(key=lambda x: x["start_date"])
            ordered_qualifications = sorted(grouped.keys(), key=str.casefold)

            for qualification in ordered_qualifications:
                entries = grouped[qualification]

                # Precompute this qualification table's full height (label row +
                # header row + one row per assessor assignment) before writing
                # anything, so we can decide up front whether it needs to start
                # on a fresh printed page.
                flat_rows = []
                for item in entries:
                    assessor_names = [a.get("name", "") for a in item.get("assessors") or [] if a.get("name")]
                    for name in (assessor_names or [""]):
                        flat_rows.append([name, item.get("assessment_center") or "", item.get("date_label") or ""])
                data_row_heights = [
                    autofit_row_height(values, [base_widths[0], base_widths[1], date_col_width])
                    for values in flat_rows
                ]
                qual_label_end_preview = name_cols[0] + 2
                qual_value_end_preview = center_cols[0] + 2
                province_label_col_preview = qual_value_end_preview + 1
                qual_row_widths_preview = [
                    sum(widths[name_cols[0] - 1:qual_label_end_preview]),
                    sum(widths[qual_label_end_preview:qual_value_end_preview]),
                    widths[province_label_col_preview - 1],
                    sum(widths[date_cols[0] - 1:date_cols[1]]),
                ]
                qual_row_height = autofit_row_height(
                    ["QUALIFICATION TITLE", qualification, province_label, province_value],
                    qual_row_widths_preview,
                    bold=True,
                )
                header_row_height = autofit_row_height(
                    headers, [base_widths[0], base_widths[1], date_col_width], bold=True
                )
                block_height = qual_row_height + header_row_height + sum(data_row_heights)

                # Only break to a fresh page if doing so would actually help:
                # we're partway down the current page, AND the block would
                # fit within a single fresh page. A block taller than a whole
                # page can't be helped by breaking (it would still overflow
                # onto further pages either way), so it's left to start right
                # where it is rather than wasting a page jumping to a new one.
                if (
                    page_used_pt > 0
                    and block_height <= usable_height_pt
                    and page_used_pt + block_height > usable_height_pt
                ):
                    sheet.row_breaks.append(Break(id=row - 1))
                    page_used_pt = 0

                # QUALIFICATION TITLE and PROVINCE/REGION share one row, split at
                # the fractions of the NAME/CENTER/DATE columns described above —
                # the same uneven split the official form uses, not a 2x2 grid.
                qual_label_end = name_cols[0] + 2  # first 3/4 of NAME OF ASSESSOR (cols 1-3)
                qual_value_end = center_cols[0] + 2  # + first 3/4 of ASSESSMENT CENTER (cols 4-7)
                province_label_col = qual_value_end + 1  # last 1/4 of ASSESSMENT CENTER (col 8)
                province_value_start = date_cols[0]  # all of DATE OF ASSESSMENT (cols 9-12)

                qual_row_cols = [
                    (name_cols[0], qual_label_end, "QUALIFICATION TITLE", label_font, left),
                    (qual_label_end + 1, qual_value_end, qualification, body_font, left),
                    (province_label_col, province_label_col, province_label, label_font, left),
                    (province_value_start, date_cols[1], province_value, body_font, left),
                ]
                for start_col, end_col, value, font, align in qual_row_cols:
                    if end_col > start_col:
                        sheet.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
                    cell = sheet.cell(row, start_col, value)
                    cell.font = font
                    cell.alignment = align
                    cell.border = thin
                    for col in range(start_col, end_col + 1):
                        sheet.cell(row, col).border = thin
                sheet.row_dimensions[row].height = qual_row_height
                row += 1

                header_cols = [
                    (name_cols[0], name_cols[1], headers[0]),
                    (center_cols[0], center_cols[1], headers[1]),
                    (date_cols[0], date_cols[1], headers[2]),
                ]
                for start_col, end_col, header in header_cols:
                    if end_col > start_col:
                        sheet.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
                    cell = sheet.cell(row, start_col, header)
                    cell.font = header_font
                    cell.alignment = center
                    cell.border = thin
                    for col in range(start_col, end_col + 1):
                        sheet.cell(row, col).border = thin
                sheet.row_dimensions[row].height = header_row_height
                row += 1

                for values, data_row_height in zip(flat_rows, data_row_heights):
                    data_cols = [
                        (name_cols[0], name_cols[1], values[0], left),
                        (center_cols[0], center_cols[1], values[1], left),
                        (date_cols[0], date_cols[1], values[2], center),
                    ]
                    for start_col, end_col, value, align in data_cols:
                        if end_col > start_col:
                            sheet.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
                        cell = sheet.cell(row, start_col, value)
                        cell.font = body_font
                        cell.alignment = align
                        cell.border = thin
                        for col in range(start_col, end_col + 1):
                            sheet.cell(row, col).border = thin
                    sheet.row_dimensions[row].height = data_row_height
                    row += 1
                row += QUAL_SPACER_ROWS  # breathing room between qualification tables
                page_used_pt += block_height + QUAL_SPACER_ROWS * 15

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
